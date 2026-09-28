"""Extract seed data (products, company, customer, sample invoice) from the legacy Excel workbook."""

import json
import sys
import warnings
from datetime import date
from pathlib import Path
from typing import Any

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
WORKBOOK = ROOT / "Invoice Generator SHJ033.xlsm"
SEED_DIR = ROOT / "backend" / "app" / "seed"

PRODUCTS_FIRST_ROW = 4
BOXES_FIRST_ITEM_ROW = 7
BOXES_FIRST_COL = 3
BOXES_LAST_COL = 102


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def to_float(value: Any) -> float | None:
    try:
        return float(value) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def extract_products(ws) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]]]:
    products: list[dict[str, Any]] = []
    selection_by_label: dict[str, dict[str, Any]] = {}
    for row in ws.iter_rows(min_row=PRODUCTS_FIRST_ROW, max_col=12, values_only=True):
        selected, code, name, material, origin, weight, qty, price, category, tariff, label, _ = row
        if not clean(code) and not clean(name):
            continue
        product = {
            "code": clean(code),
            "name": clean(name),
            "material": clean(material),
            "origin": clean(origin),
            "weight_g": to_float(weight),
            "unit_price": to_float(price),
            "category": clean(category),
            "tariff_code": clean(tariff),
        }
        products.append(product)
        if selected == 1 and to_float(qty):
            selection_by_label[clean(label)] = {"index": len(products) - 1, "qty": int(to_float(qty))}
    return products, selection_by_label


def extract_boxes(ws) -> list[dict[str, Any]]:
    """Box dimensions/weights only: the workbook's box contents belong to an older order and do not reconcile."""
    boxes: list[dict[str, Any]] = []
    for col in range(BOXES_FIRST_COL, BOXES_LAST_COL + 1):
        weight = to_float(ws.cell(4, col).value)
        if not weight:
            continue
        boxes.append(
            {
                "length_cm": to_float(ws.cell(1, col).value),
                "width_cm": to_float(ws.cell(2, col).value),
                "height_cm": to_float(ws.cell(3, col).value),
                "gross_weight_kg": weight,
                "items": [],
            }
        )
    return boxes


def distribute_lines(boxes: list[dict[str, Any]], lines: list[dict[str, Any]], products: list[dict[str, Any]]) -> None:
    """Fill boxes sequentially so each box's net weight is proportional to its gross weight."""
    unit_kg = [(products[line["product_index"]]["weight_g"] or 0) / 1000 for line in lines]
    total_net = sum(kg * line["qty"] for kg, line in zip(unit_kg, lines))
    total_gross = sum(box["gross_weight_kg"] for box in boxes)
    remaining = [line["qty"] for line in lines]
    line_idx = 0
    for box_idx, box in enumerate(boxes):
        is_last = box_idx == len(boxes) - 1
        capacity = total_net * box["gross_weight_kg"] / total_gross
        while line_idx < len(lines) and (capacity > 0 or is_last):
            kg = unit_kg[line_idx] or 1e-9
            fits = int(capacity / kg) if box["items"] else max(1, int(capacity / kg))
            take = remaining[line_idx] if is_last else min(remaining[line_idx], fits)
            if take == 0:
                break
            box["items"].append({"product_index": lines[line_idx]["product_index"], "qty": take})
            capacity -= take * kg
            remaining[line_idx] -= take
            if remaining[line_idx] == 0:
                line_idx += 1


def main() -> None:
    warnings.simplefilter("ignore")
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    products, selection = extract_products(wb["Products"])
    boxes = extract_boxes(wb["Boxes"])
    invoice_ws = wb["Invoice"]

    company = {
        "name": clean(invoice_ws["B6"].value),
        "address": "\n".join(clean(invoice_ws[c].value) for c in ("B7", "B8", "B9")),
        "phone": clean(invoice_ws["C11"].value),
        "fax": "",
        "place": clean(invoice_ws["L4"].value),
        "shipped_per": clean(invoice_ws["L10"].value),
        "payment_terms": " ".join(clean(invoice_ws["J13"].value).split()),
        "currency_label": "JPY日本円",
        "currency_symbol": "¥",
        "trade_term_label": "E.X.W.JAPAN",
        "invoice_prefix": "SHJ",
        "next_invoice_number": 34,
    }
    customer = {
        "name": clean(invoice_ws["B14"].value),
        "address": "\n".join(clean(invoice_ws[c].value) for c in ("B15", "B16", "B17", "B18")),
        "phone": clean(invoice_ws["C20"].value),
        "fax": "",
    }
    lines = [{"product_index": s["index"], "qty": s["qty"]} for s in selection.values()]
    distribute_lines(boxes, lines, products)
    sample_invoice = {
        "number": "SHJ033",
        "invoice_date": date(2026, 9, 28).isoformat(),
        "lines": lines,
        "boxes": boxes,
    }

    SEED_DIR.mkdir(parents=True, exist_ok=True)
    payload = {"company": company, "customer": customer, "products": products, "sample_invoice": sample_invoice}
    (SEED_DIR / "seed.json").write_text(json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"products={len(products)} selected={len(selection)} boxes={len(boxes)}", file=sys.stderr)


if __name__ == "__main__":
    main()
