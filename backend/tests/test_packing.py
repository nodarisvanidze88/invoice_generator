from collections import Counter
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.schemas import InvoiceLine
from app.services.packing import BoxSpec, auto_pack, category_densities, estimate_dims

DIVISOR = 5000
LARGE = BoxSpec(1, "Large", (60, 40, 40), 1.0, 30)
SMALL = BoxSpec(2, "Small", (30, 20, 20), 0.5, 10)


def line(qty: int, dims: tuple[float, float, float] | None = (10, 10, 10), weight_g: float = 100, **extra) -> InvoiceLine:
    length, width, height = dims or (None, None, None)
    return InvoiceLine(name="Item", qty=qty, weight_g=weight_g, length_cm=length, width_cm=width, height_cm=height, **extra)


def packed_quantities(result) -> Counter[str]:
    counts: Counter[str] = Counter()
    for box in result.boxes:
        for block in box.blocks:
            counts[block.line_id] += block.count
    return counts


def assert_valid_geometry(result) -> None:
    for box in result.boxes:
        for index, block in enumerate(box.blocks):
            assert all(block.end[a] <= box.spec.inner[a] + 1e-6 for a in range(3))
            assert all(not block.intersects(other.origin, other.end) for other in box.blocks[index + 1 :])
        assert box.gross_kg <= box.spec.max_kg + 1e-6


def test_exact_fill_uses_one_box() -> None:
    cubes = line(qty=96)
    result = auto_pack([cubes], [LARGE], DIVISOR, {})
    assert len(result.boxes) == 1
    assert packed_quantities(result)[cubes.id] == 96
    assert result.boxes[0].used_volume == pytest.approx(LARGE.volume)
    assert_valid_geometry(result)


def test_items_rotate_to_fit() -> None:
    plank = line(qty=1, dims=(5, 5, 55))
    result = auto_pack([plank], [LARGE], DIVISOR, {})
    assert packed_quantities(result)[plank.id] == 1
    assert_valid_geometry(result)


def test_weight_limit_splits_boxes() -> None:
    heavy = line(qty=10, weight_g=5000)
    result = auto_pack([heavy], [LARGE], DIVISOR, {})
    assert [sum(b.count for b in box.blocks) for box in result.boxes] == [5, 5]
    assert_valid_geometry(result)


def test_small_order_picks_box_with_lower_chargeable_weight() -> None:
    few = line(qty=4)
    result = auto_pack([few], [LARGE, SMALL], DIVISOR, {})
    assert [box.spec.name for box in result.boxes] == ["Small"]


def test_mixed_lines_all_packed_without_overlap() -> None:
    lines = [line(qty=37, dims=(12, 7, 4)), line(qty=11, dims=(25, 18, 9), weight_g=900), line(qty=150, dims=(3, 3, 8), weight_g=20)]
    result = auto_pack(lines, [LARGE, SMALL], DIVISOR, {})
    counts = packed_quantities(result)
    assert all(counts[item.id] == item.qty for item in lines)
    assert not result.unpacked
    assert_valid_geometry(result)


def test_oversized_and_overweight_items_reported() -> None:
    too_big = line(qty=2, dims=(100, 10, 10))
    too_heavy = line(qty=1, weight_g=50_000)
    result = auto_pack([too_big, too_heavy], [LARGE], DIVISOR, {})
    reasons = {u.line_id: u.reason for u in result.unpacked}
    assert reasons == {too_big.id: "larger than every box", too_heavy.id: "heavier than every box allows"}


def test_missing_dimensions_estimated_from_category_density() -> None:
    densities = category_densities([("Candy", 1000, 10, 10, 10), ("Candy", 500, 10, 10, 5)])
    assert densities["Candy"] == pytest.approx(1.0)
    dims, estimated = estimate_dims(line(qty=1, dims=None, weight_g=1000, category="Candy"), densities)
    assert estimated and dims == (10.0, 10.0, 10.0)


def test_auto_pack_endpoint_and_box_types(auth_client: TestClient) -> None:
    box_type = auth_client.post(
        "/api/box-types", json={"name": "Large", "length_cm": 60, "width_cm": 40, "height_cm": 40, "tare_weight_kg": 1, "max_weight_kg": 30}
    ).json()
    assert [bt["name"] for bt in auth_client.get("/api/box-types").json()] == ["Large"]

    item = line(qty=20, weight_g=1000).model_dump()
    result = auth_client.post("/api/packing/auto", json={"lines": [item]}).json()
    assert result["volumetric_divisor"] == 5000
    [box] = result["boxes"]
    assert box["box_type_id"] == box_type["id"]
    assert box["items"] == [{"line_id": item["id"], "qty": 20}]
    assert box["gross_weight_kg"] == 21.0
    assert sum(p["nx"] * p["ny"] * p["nz"] for p in box["placements"]) == 20


def test_auto_pack_requires_box_types(auth_client: TestClient) -> None:
    response = auth_client.post("/api/packing/auto", json={"lines": [line(qty=1).model_dump()]})
    assert response.status_code == 422


def test_product_excel_round_trip_updates_dimensions(auth_client: TestClient, make_product) -> None:
    product = make_product()
    exported = load_workbook(BytesIO(auth_client.get("/api/products/export.xlsx").content))
    sheet = exported.active
    header = [cell.value for cell in sheet[1]]
    sheet.cell(row=2, column=header.index("Length (cm)") + 1, value=12.5)
    sheet.cell(row=2, column=header.index("Width (cm)") + 1, value=8)
    sheet.cell(row=2, column=header.index("Height (cm)") + 1, value=3)
    sheet.append([None, "4999", "Brand new product", "", "Japan", 50, 300, "Toy", "95030070", 5, 5, 5])
    sheet.append([None, "5000", "", "", "", "abc"])
    buffer = BytesIO()
    exported.save(buffer)

    summary = auth_client.post("/api/products/import.xlsx?mode=upsert", files={"file": ("p.xlsx", buffer.getvalue())}).json()
    assert (summary["updated"], summary["created"], summary["skipped_invalid"]) == (1, 1, 1)
    updated = next(p for p in auth_client.get("/api/products").json() if p["id"] == product["id"])
    assert (updated["length_cm"], updated["width_cm"], updated["height_cm"]) == (12.5, 8, 3)
    assert updated["unit_price"] == product["unit_price"]


def test_product_import_rejects_unknown_layout(auth_client: TestClient) -> None:
    workbook = Workbook()
    workbook.active.append(["foo", "bar"])
    buffer = BytesIO()
    workbook.save(buffer)
    response = auth_client.post("/api/products/import.xlsx", files={"file": ("x.xlsx", buffer.getvalue())})
    assert response.status_code == 422


def test_upright_items_keep_height_vertical() -> None:
    bottle = line(qty=6, dims=(8, 8, 30), keep_upright=True)
    result = auto_pack([bottle], [LARGE], DIVISOR, {})
    assert packed_quantities(result)[bottle.id] == 6
    assert all(block.unit[2] == 30 for box in result.boxes for block in box.blocks)


def test_upright_item_taller_than_every_box_is_reported() -> None:
    poster = line(qty=1, dims=(5, 5, 55), keep_upright=True)
    result = auto_pack([poster], [LARGE], DIVISOR, {})
    assert [(u.line_id, u.reason) for u in result.unpacked] == [(poster.id, "too tall to stand upright in any box")]


def test_keep_upright_round_trips_through_excel(auth_client: TestClient, make_product) -> None:
    product = make_product()
    workbook = load_workbook(BytesIO(auth_client.get("/api/products/export.xlsx").content))
    sheet = workbook.active
    header = [cell.value for cell in sheet[1]]
    column = header.index("Keep upright") + 1
    assert sheet.cell(row=2, column=column).value == "No"
    sheet.cell(row=2, column=column, value="Yes")
    buffer = BytesIO()
    workbook.save(buffer)
    auth_client.post("/api/products/import.xlsx?mode=upsert", files={"file": ("p.xlsx", buffer.getvalue())})
    assert next(p for p in auth_client.get("/api/products").json() if p["id"] == product["id"])["keep_upright"] is True


def test_gap_fillers_are_not_scattered_across_boxes() -> None:
    heavy = line(qty=40, dims=(20, 20, 20), weight_g=4000)
    bottles = line(qty=12, dims=(6, 6, 18), weight_g=300, keep_upright=True)
    result = auto_pack([heavy, bottles], [LARGE], DIVISOR, {})
    boxes_with_bottles = [box for box in result.boxes if any(b.line_id == bottles.id for b in box.blocks)]
    assert len(boxes_with_bottles) == 1
    assert packed_quantities(result)[bottles.id] == 12
    assert_valid_geometry(result)


def test_layout_endpoint_arranges_existing_box_contents(auth_client: TestClient) -> None:
    fits = line(qty=24)
    overflow = line(qty=500, dims=(10, 10, 10))
    body = {
        "length_cm": 40, "width_cm": 30, "height_cm": 20,
        "items": [{"line_id": fits.id, "qty": 24}, {"line_id": overflow.id, "qty": 500}],
        "lines": [fits.model_dump(), overflow.model_dump()],
    }
    result = auth_client.post("/api/packing/layout", json=body).json()
    placed = Counter()
    for p in result["placements"]:
        placed[p["line_id"]] += p["nx"] * p["ny"] * p["nz"]
    assert sum(placed.values()) == 24
    assert sum(u["qty"] for u in result["unplaced"]) == 500
    assert result["estimated_line_ids"] == []
