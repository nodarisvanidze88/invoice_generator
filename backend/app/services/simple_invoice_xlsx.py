"""Simplified invoice as an Excel workbook: same layout as the PDF, plain numeric values (no formulas) for copy/paste."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas import Party
from app.services.simple_invoice import SimpleInvoice

ACCENT = "2F5496"
HEADER_FILL = PatternFill("solid", fgColor=ACCENT)
WHITE_BOLD = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="333333")
THICK = Side(style="medium", color="000000")
COLUMN_WIDTHS = {"A": 48, "B": 14, "C": 12, "D": 14, "E": 20}
WEIGHT_FORMAT = '0.00" Kg"'
QTY_FORMAT = "#,##0"
PRICE_FORMAT = "#,##0.00"
TABLE_HEADERS = ("DESCRIPTIONS [HS Code]", "Weight (Kg)", "Quantity", "Unit Price", "TOTAL PRICE")


def _party_rows(party: Party) -> list[str]:
    lines = [f"Company: {party.name}"]
    address = [line.strip() for line in party.address.splitlines() if line.strip()]
    lines += [f"Address: {address[0]}", *address[1:]] if address else ["Address:"]
    if party.phone:
        lines.append(f"TEL: {party.phone}")
    return lines


def _write_header(sheet: Worksheet, invoice: SimpleInvoice) -> int:
    sheet["A1"] = invoice.company_name.upper()
    sheet["A1"].font = Font(size=16)
    sheet["E1"] = "INVOICE"
    sheet["E1"].font = Font(size=16, bold=True, color=ACCENT)
    sheet["E1"].alignment = Alignment(horizontal="right")
    sheet["A2"] = f"Address: {invoice.company_address}"
    sheet["A3"] = f"TEL: {invoice.company_phone}" if invoice.company_phone else ""

    sheet["A5"], sheet["B5"] = "INVOICE #:", invoice.number
    sheet["D5"], sheet["E5"] = "INVOICE DATE:", invoice.invoice_date_label
    for cell in ("A5", "D5"):
        sheet[cell].font = BOLD
    for column in "ABCDE":
        sheet[f"{column}5"].border = Border(top=THIN, bottom=THIN)

    sheet["A7"], sheet["D7"] = "SHIPPING TO", "BILL TO"
    sheet.merge_cells("D7:E7")
    for cell in ("A7", "B7", "D7", "E7"):
        sheet[cell].fill = HEADER_FILL
        sheet[cell].font = WHITE_BOLD
    shipping, billing = _party_rows(invoice.shipping_to), _party_rows(invoice.bill_to)
    row = 8
    for index in range(max(len(shipping), len(billing))):
        sheet.cell(row=row + index, column=1, value=shipping[index] if index < len(shipping) else None)
        sheet.cell(row=row + index, column=4, value=billing[index] if index < len(billing) else None)
    return row + max(len(shipping), len(billing)) + 1


def _write_table(sheet: Worksheet, invoice: SimpleInvoice, start_row: int) -> None:
    headers = list(TABLE_HEADERS)
    headers[-1] = f"TOTAL PRICE ({invoice.currency_code})"
    for column, title in enumerate(headers, start=1):
        cell = sheet.cell(row=start_row, column=column, value=title)
        cell.font = BOLD
        cell.alignment = Alignment(horizontal="left" if column == 1 else "center")
        cell.border = Border(bottom=THIN)

    first_data_row = start_row + 1
    for offset, item in enumerate(invoice.rows):
        row = first_data_row + offset
        sheet.cell(row=row, column=1, value=item.label).font = BOLD
        sheet.cell(row=row, column=2, value=round(item.net_weight_kg, 3)).number_format = WEIGHT_FORMAT
        sheet.cell(row=row, column=3, value=item.qty).number_format = QTY_FORMAT
        sheet.cell(row=row, column=4, value=round(item.unit_price, 4)).number_format = PRICE_FORMAT
        amount = sheet.cell(row=row, column=5, value=round(item.amount, 4))
        amount.number_format = PRICE_FORMAT
        amount.font = BOLD

    last_data_row = first_data_row + max(len(invoice.rows), 1) - 1
    total_row = last_data_row + 1
    sheet.cell(row=total_row, column=1, value=f"Total amount ({invoice.currency_code})")
    sheet.cell(row=total_row, column=2, value=round(invoice.total_weight_kg, 3)).number_format = WEIGHT_FORMAT
    sheet.cell(row=total_row, column=3, value=invoice.total_qty).number_format = QTY_FORMAT
    sheet.cell(row=total_row, column=5, value=round(invoice.total_amount, 2)).number_format = PRICE_FORMAT
    for column in range(1, 6):
        cell = sheet.cell(row=total_row, column=column)
        cell.font = BOLD
        cell.border = Border(top=THICK)
    for row in range(first_data_row, total_row + 1):
        for column in (2, 3, 4):
            sheet.cell(row=row, column=column).alignment = Alignment(horizontal="center")


def simple_invoice_to_xlsx(invoice: SimpleInvoice) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f"Invoice {invoice.number}"[:31]
    table_start = _write_header(sheet, invoice)
    _write_table(sheet, invoice, table_start)
    for column, width in COLUMN_WIDTHS.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = sheet.cell(row=table_start + 1, column=1)
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = f"{table_start}:{table_start}"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
