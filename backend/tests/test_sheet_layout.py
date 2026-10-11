import json
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.services.sheet_layout import inspect_workbook, match_headers, read_table


def multi_sheet_workbook() -> bytes:
    workbook = Workbook()
    notes = workbook.active
    notes.title = "Notes"
    notes.append(["Product Name", "Price"])
    notes.append(["", ""])
    order = workbook.create_sheet("Order")
    order.append(["Order summary for September"])
    order.append([])
    order.append(["", "商品名", "重量 (kg)", "単価", "カテゴリ", "HS Code", "Comment"])
    order.append([1, "Matcha KitKat", 0.25, 598, "Sugar-based Candy", "17049099", "fragile"])
    order.append([2, "Daruma Figure", 1.2, 1250, "Plastic Toy", "95030070", ""])
    for row in order.iter_rows(min_row=4, max_row=5, min_col=1, max_col=1):
        row[0].value = {4: "4902201178921", 5: "4580123450011"}[row[0].row]
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_header_aliases_prefer_exact_matches() -> None:
    mapping, in_kg = match_headers(["Product Name", "Product concatenated", "Product Category", "Product Category concatenated", "Weight(g)"])
    assert mapping == {"name": 0, "category": 2, "weight_g": 4}
    assert in_kg is False


def test_inspect_finds_data_sheet_header_row_and_unlabelled_jan_column() -> None:
    result = inspect_workbook(multi_sheet_workbook())
    assert result.layout.sheet == "Order"
    assert result.layout.header_row == 3
    assert result.layout.mapping == {"name": 1, "weight_g": 2, "unit_price": 3, "category": 4, "tariff_code": 5, "code": 0}
    assert result.layout.weight_in_kg is True
    assert result.data_rows == 2
    assert result.unmapped_headers == ["G · Comment"]


def test_read_table_converts_kg_to_grams() -> None:
    content = multi_sheet_workbook()
    rows = read_table(content, inspect_workbook(content).layout)
    assert [(number, values["code"], values["name"], values["weight_g"]) for number, values in rows] == [
        (4, "4902201178921", "Matcha KitKat", 250.0),
        (5, "4580123450011", "Daruma Figure", 1200.0),
    ]


def test_inspect_endpoint_and_import_with_user_mapping(auth_client: TestClient) -> None:
    files = {"file": ("order.xlsx", multi_sheet_workbook())}
    inspection = auth_client.post("/api/products/import.xlsx/inspect", files=files).json()
    assert [s["name"] for s in inspection["sheets"] if s["recommended"]] == ["Order"]
    assert inspection["columns"][1]["samples"] == ["Matcha KitKat", "Daruma Figure"]

    layout = inspection["layout"] | {"mapping": inspection["layout"]["mapping"] | {"material": 6}}
    preview = auth_client.post("/api/products/import.xlsx/preview", files=files, data={"layout": json.dumps(layout)}).json()
    assert [r["name"] for r in preview["new"]] == ["Matcha KitKat", "Daruma Figure"]

    summary = auth_client.post("/api/products/import.xlsx", files=files, data={"layout": json.dumps(layout)}).json()
    assert summary["created"] == 2
    kitkat = next(p for p in auth_client.get("/api/products").json() if p["name"] == "Matcha KitKat")
    assert (kitkat["code"], kitkat["weight_g"], kitkat["material"]) == ("4902201178921", 250, "fragile")


def test_mapping_without_name_is_rejected(auth_client: TestClient) -> None:
    layout = {"sheet": "Order", "header_row": 3, "mapping": {"code": 0}}
    response = auth_client.post(
        "/api/products/import.xlsx/preview", files={"file": ("o.xlsx", multi_sheet_workbook())}, data={"layout": json.dumps(layout)}
    )
    assert response.status_code == 422
    assert "Product Name" in response.json()["detail"]


def test_choosing_another_sheet_and_header_row(auth_client: TestClient) -> None:
    files = {"file": ("order.xlsx", multi_sheet_workbook())}
    inspection = auth_client.post("/api/products/import.xlsx/inspect", files=files, data={"sheet": "Notes", "header_row": "1"}).json()
    assert inspection["layout"]["sheet"] == "Notes"
    assert inspection["layout"]["mapping"] == {"name": 0, "unit_price": 1}
    assert inspection["data_rows"] == 0


def test_import_template_matches_export_headers(auth_client: TestClient) -> None:
    template = load_workbook(BytesIO(auth_client.get("/api/products/import-template.xlsx").content)).active
    export = load_workbook(BytesIO(auth_client.get("/api/products/export.xlsx").content)).active
    assert [c.value for c in template[1]] == [c.value for c in export[1]]
    assert template.max_row == 3


def test_not_an_excel_file(auth_client: TestClient) -> None:
    response = auth_client.post("/api/products/import.xlsx/inspect", files={"file": ("x.xlsx", b"hello")})
    assert response.status_code == 422
