from io import BytesIO

from fastapi.testclient import TestClient
from pypdf import PdfReader

PDF_MAGIC = b"%PDF"


def test_protected_routes_require_login(client: TestClient) -> None:
    assert client.get("/api/products").status_code == 401
    assert client.post("/api/auth/login", json={"password": "wrong"}).status_code == 401
    assert client.get("/api/auth/me").json() == {"authenticated": False}


def test_logout_clears_session(auth_client: TestClient) -> None:
    auth_client.post("/api/auth/logout")
    assert auth_client.get("/api/products").status_code == 401


def test_product_crud(auth_client: TestClient, make_product) -> None:
    product = make_product()
    updated = auth_client.put(f"/api/products/{product['id']}", json={**product, "unit_price": 140}).json()
    assert updated["unit_price"] == 140
    assert [p["id"] for p in auth_client.get("/api/products").json()] == [product["id"]]
    assert auth_client.delete(f"/api/products/{product['id']}").status_code == 204
    assert auth_client.get("/api/products").json() == []


def test_product_validation(auth_client: TestClient) -> None:
    response = auth_client.post("/api/products", json={"name": "", "unit_price": -1})
    assert response.status_code == 422


def test_bulk_delete_and_export(auth_client: TestClient, make_product) -> None:
    ids = [make_product(code=str(i))["id"] for i in range(3)]
    export = auth_client.get("/api/products/export.xlsx")
    assert export.status_code == 200
    assert export.content[:2] == b"PK"
    assert auth_client.post("/api/products/bulk-delete", json={"ids": ids[:2]}).status_code == 204
    assert [p["id"] for p in auth_client.get("/api/products").json()] == [ids[2]]


def test_create_invoice_snapshots_products_and_increments_number(auth_client: TestClient, make_product) -> None:
    auth_client.put("/api/company", json={"name": "Shinon Japan", "invoice_prefix": "SHJ", "next_invoice_number": 34})
    customer = auth_client.post("/api/customers", json={"name": "New Japan Trading AS", "address": "Oslo"}).json()
    product = make_product()

    invoice = auth_client.post(
        "/api/invoices", json={"items": [{"product_id": product["id"], "qty": 36}], "customer_id": customer["id"]}
    ).json()

    assert invoice["number"] == "SHJ034"
    assert invoice["addressee"]["name"] == "New Japan Trading AS"
    assert invoice["sender"]["name"] == "Shinon Japan"
    assert invoice["lines"][0]["qty"] == 36
    assert invoice["totals"]["amount"] == 36 * 132.5
    assert auth_client.get("/api/invoices/next-number").json() == {"number": "SHJ035"}

    auth_client.put(f"/api/products/{product['id']}", json={**product, "unit_price": 999})
    assert auth_client.get(f"/api/invoices/{invoice['id']}").json()["lines"][0]["unit_price"] == 132.5


def test_duplicate_invoice_number_rejected(auth_client: TestClient, make_product) -> None:
    product = make_product()
    body = {"items": [{"product_id": product["id"], "qty": 1}], "number": "X-1"}
    assert auth_client.post("/api/invoices", json=body).status_code == 201
    assert auth_client.post("/api/invoices", json=body).status_code == 409


def test_update_invoice_with_boxes_and_render_pdfs(auth_client: TestClient, make_product) -> None:
    product = make_product()
    invoice = auth_client.post("/api/invoices", json={"items": [{"product_id": product["id"], "qty": 10}]}).json()
    line_id = invoice["lines"][0]["id"]
    doc = {k: v for k, v in invoice.items() if k not in {"id", "created_at", "updated_at", "totals"}}
    doc["boxes"] = [{"length_cm": 64, "width_cm": 43, "height_cm": 52, "gross_weight_kg": 2.5, "items": [{"line_id": line_id, "qty": 10}]}]

    saved = auth_client.put(f"/api/invoices/{invoice['id']}", json=doc).json()
    assert saved["totals"]["pieces"] == 1
    assert saved["totals"]["mismatched_lines"] == 0

    for kind in ("invoice", "packing"):
        response = auth_client.get(f"/api/invoices/{invoice['id']}/{kind}.pdf")
        assert response.status_code == 200
        assert response.content.startswith(PDF_MAGIC)
        text = "".join(page.extract_text() for page in PdfReader(BytesIO(response.content)).pages)
        assert invoice["number"] in text


def test_update_invoice_rejects_dangling_box_items(auth_client: TestClient, make_product) -> None:
    product = make_product()
    invoice = auth_client.post("/api/invoices", json={"items": [{"product_id": product["id"], "qty": 1}]}).json()
    doc = {k: v for k, v in invoice.items() if k not in {"id", "created_at", "updated_at", "totals"}}
    doc["boxes"] = [{"items": [{"line_id": "ghost", "qty": 1}]}]
    assert auth_client.put(f"/api/invoices/{invoice['id']}", json=doc).status_code == 422


def test_duplicate_and_delete_invoice(auth_client: TestClient, make_product) -> None:
    product = make_product()
    original = auth_client.post("/api/invoices", json={"items": [{"product_id": product["id"], "qty": 2}]}).json()
    copy = auth_client.post(f"/api/invoices/{original['id']}/duplicate").json()
    assert copy["number"] != original["number"]
    assert copy["lines"][0]["id"] != original["lines"][0]["id"]
    assert auth_client.delete(f"/api/invoices/{original['id']}").status_code == 204
    assert [i["id"] for i in auth_client.get("/api/invoices").json()] == [copy["id"]]


def test_csv_import_preview_and_apply(auth_client: TestClient, make_product) -> None:
    existing = make_product(code="4901234567890 / 4901234567891")
    csv_text = "注文日,商品,JAN,数量,単価\n2026/09/01,Candy,4901234567891,24,150\n2026/09/01,New Toy,4580000000001,5,900\n"

    preview = auth_client.post(
        "/api/imports/csv/preview", files={"file": ("order.csv", csv_text.encode("cp932"), "text/csv")}
    ).json()
    assert preview["encoding"] == "cp932"
    assert (preview["matched_count"], preview["new_count"]) == (1, 1)
    matched, new = preview["rows"]
    assert matched["product_id"] == existing["id"]

    rows = [
        {**matched["current"], "product_id": existing["id"], "unit_price": 150, "qty": 24},
        {"code": new["code"], "name": "New Toy", "unit_price": 900, "qty": 5, "category": "Plastic Toy"},
    ]
    result = auth_client.post("/api/imports/csv/apply", json={"rows": rows, "update_prices": True}).json()
    assert (result["created"], result["updated"]) == (1, 1)
    assert len(result["selection"]) == 2
    prices = {p["name"]: p["unit_price"] for p in auth_client.get("/api/products").json()}
    assert prices == {"Pineapple Candy 110g": 150, "New Toy": 900}


def test_csv_without_jan_column_is_rejected(auth_client: TestClient) -> None:
    response = auth_client.post("/api/imports/csv/preview", files={"file": ("x.csv", b"name,qty\nA,1\n", "text/csv")})
    assert response.status_code == 422


def test_customer_bill_to_flows_into_invoice(auth_client: TestClient, make_product) -> None:
    billing = {"name": "Neo Tokyo Norway AS", "address": "Karl Johans gate 7\n0154 Oslo", "phone": "+47 210 95 990", "fax": ""}
    customer = auth_client.post(
        "/api/customers", json={"name": "Neo Tokyo Norway AS", "address": "Mølleparken 4\n0459 Oslo", "bill_to": billing}
    ).json()
    assert customer["bill_to"]["address"].startswith("Karl Johans")
    product = make_product()
    invoice = auth_client.post(
        "/api/invoices", json={"items": [{"product_id": product["id"], "qty": 3}], "customer_id": customer["id"]}
    ).json()
    assert invoice["addressee"]["address"].startswith("Mølleparken")
    assert invoice["bill_to"] == billing


def test_simplified_invoice_pdf_and_excel(auth_client: TestClient, make_product) -> None:
    from openpyxl import load_workbook

    first = make_product(code="1", unit_price=100, weight_g=500)
    second = make_product(code="2", unit_price=50, weight_g=250)
    other = make_product(code="3", category="Hair Oil", tariff_code="33059000", unit_price=847.9496, weight_g=80)
    items = [{"product_id": first["id"], "qty": 2}, {"product_id": second["id"], "qty": 6}, {"product_id": other["id"], "qty": 36}]
    invoice = auth_client.post("/api/invoices", json={"items": items}).json()

    pdf = auth_client.get(f"/api/invoices/{invoice['id']}/simple.pdf")
    assert pdf.content.startswith(PDF_MAGIC)
    text = "".join(page.extract_text() for page in PdfReader(BytesIO(pdf.content)).pages)
    assert "Sugar-based Candy [17049099]" in text and "BILL TO" in text

    xlsx = auth_client.get(f"/api/invoices/{invoice['id']}/simple.xlsx")
    assert xlsx.status_code == 200
    rows = [r for r in load_workbook(BytesIO(xlsx.content)).active.iter_rows(values_only=True) if r[0]]
    header = next(i for i, r in enumerate(rows) if r[0] == "DESCRIPTIONS [HS Code]")
    candy, oil, total = rows[header + 1], rows[header + 2], rows[header + 3]
    assert candy[:5] == ("Sugar-based Candy [17049099]", 2.5, 8, 62.5, 500)
    assert oil[0] == "Hair Oil [33059000]" and round(oil[3], 2) == 847.95 and round(oil[4], 2) == 30526.19
    assert total[0] == "Total amount (JPY)" and total[2] == 44 and round(total[4], 2) == 31026.19


def test_full_invoice_excel(auth_client: TestClient, make_product) -> None:
    from openpyxl import load_workbook

    candy = make_product(code="1", unit_price=100, weight_g=500)
    oil = make_product(code="2", category="Hair Oil", tariff_code="33059000", unit_price=847.9496, weight_g=80, origin="China")
    items = [{"product_id": candy["id"], "qty": 2}, {"product_id": oil["id"], "qty": 36}]
    invoice = auth_client.post("/api/invoices", json={"items": items}).json()

    response = auth_client.get(f"/api/invoices/{invoice['id']}/invoice.xlsx")
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content)).active
    rows = [r for r in sheet.iter_rows(values_only=True) if any(v is not None for v in r)]
    assert rows[0][0] == f"INVOICE N {invoice['number']}"
    labels = [r[0] for r in rows]
    group = labels.index("Sugar-based Candy [17049099]")
    assert rows[group][1:5] == (1.0, 2, None, 200)
    assert rows[group + 1][0] == "1 Pineapple Candy 110g [Japan]" and rows[group + 1][3] == 100
    total = rows[labels.index("総合計 (Total)")]
    assert total[2] == 38 and round(total[4], 2) == round(200 + 36 * 847.9496, 2)
    assert rows[labels.index("原産国(Country of Origin)")][1] == "Japan/China"
