from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models import Product

COLUMNS: list[tuple[str, str, int]] = [
    ("ID", "id", 8),
    ("JAN / Code", "code", 28),
    ("Product Name", "name", 50),
    ("Material", "material", 18),
    ("Origin", "origin", 12),
    ("Weight(g)", "weight_g", 11),
    ("Unit Price", "unit_price", 12),
    ("Product Category", "category", 28),
    ("Tariff Code", "tariff_code", 14),
    ("Length (cm)", "length_cm", 12),
    ("Width (cm)", "width_cm", 12),
    ("Height (cm)", "height_cm", 12),
    ("Keep upright", "keep_upright", 13),
]
HEADER_FILL = PatternFill("solid", fgColor="1E293B")


def _cell(value: object) -> object:
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return value


def products_to_xlsx(products: list[Product]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append([title for title, _, _ in COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    for product in products:
        sheet.append([_cell(getattr(product, attr)) for _, attr, _ in COLUMNS])
    for index, (_, _, width) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
