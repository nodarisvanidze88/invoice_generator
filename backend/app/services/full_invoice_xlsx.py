"""Full (customs) invoice as an Excel workbook: category groups with subtotals, every line, totals and footer.

Plain numeric values only (no formulas) so rows can be copied into other workbooks.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas import InvoiceDocument, Party
from app.services.invoice_calc import compute_totals, group_lines

BOLD = Font(bold=True)
TITLE = Font(bold=True, size=16)
THIN = Side(style="thin", color="333333")
MEDIUM = Side(style="medium", color="000000")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
GROUP_FILL = PatternFill("solid", fgColor="F2F4F7")
TOTAL_FILL = PatternFill("solid", fgColor="E8EBF0")
WRAP_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
COLUMN_WIDTHS = {"A": 70, "B": 14, "C": 12, "D": 16, "E": 18}
WEIGHT_FORMAT = '0.00" Kg"'
QTY_FORMAT = "#,##0"
PRICE_FORMAT = "#,##0.00##"
AMOUNT_FORMAT = "#,##0.00"
LINE_INDENT = 2
CHECKED, UNCHECKED = "☑", "☐"


def _party_lines(label: str, party: Party) -> list[tuple[str, bool]]:
    lines: list[tuple[str, bool]] = [(label, True), (party.name, True)]
    lines += [(line, False) for line in party.address.splitlines() if line.strip()]
    lines += [(f"TEL  {party.phone}", False), (f"FAX  {party.fax}", False)]
    return lines


def _write_header(sheet: Worksheet, doc: InvoiceDocument) -> int:
    sheet["A1"] = f"INVOICE N {doc.number}"
    sheet["A1"].font = TITLE
    sheet["C2"], sheet["E2"] = "インボイス作成日(Date) :", doc.invoice_date.strftime("%d-%b-%Y")
    sheet["C3"], sheet["E3"] = "作成地(Place):", doc.place
    for cell in ("C2", "C3", "E2", "E3"):
        sheet[cell].font = BOLD

    left = _party_lines("ご依頼主 (Sender) FROM:", doc.sender) + [("", False)]
    left += _party_lines("お届け先 (Addressee) TO:", doc.addressee)
    commercial, free = (CHECKED, UNCHECKED) if doc.commercial_value else (UNCHECKED, CHECKED)
    right: list[tuple[str, bool]] = [
        ("発送番号 (Shipment No.):", True),
        (doc.number, True),
        ("", False),
        (f"送達手段 (Shipped Per): {doc.shipped_per}", True),
        ("", False),
        ("支払い条件(Terms of Payment):", True),
        (doc.payment_terms, True),
        ("", False),
        ("備考 (Remarks):", True),
        (f"{commercial} 有償 (Commercial value)", True),
        (f"{free} 無償 (No Commercial value)", True),
        ("     贈物 (Gift)   商品見本 (Sample)   その他 (Other)", False),
        *[(line, False) for line in doc.remarks.splitlines()[:3]],
    ]
    start = 5
    for offset in range(max(len(left), len(right))):
        row = start + offset
        for column, entries in ((1, left), (3, right)):
            if offset < len(entries):
                text, bold = entries[offset]
                cell = sheet.cell(row=row, column=column, value=text or None)
                cell.font = Font(bold=bold)
    return start + max(len(left), len(right)) + 1


def _write_table(sheet: Worksheet, doc: InvoiceDocument, currency_label: str, start: int) -> int:
    headers = (
        "内容品の記載\n(Description)(Tariff)",
        "正味重量\n(Net Weight)\nKg",
        "数量\n(Quantity)",
        f"単価\n(Unit Price)\n通貨(Currency)\n{currency_label}",
        "合計額\n(Total Amount)",
    )
    for column, title in enumerate(headers, start=1):
        cell = sheet.cell(row=start, column=column, value=title)
        cell.font = BOLD
        cell.alignment = WRAP_CENTER
        cell.border = BOX
    sheet.row_dimensions[start].height = 62

    row = start + 1
    for group in group_lines(doc.lines):
        values = (group.label, round(group.net_weight_kg, 3), group.qty, None, round(group.amount, 2))
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.font = BOLD
            cell.fill = GROUP_FILL
        _number_formats(sheet, row)
        row += 1
        for line in group.lines:
            sheet.cell(row=row, column=1, value=line.description).alignment = Alignment(indent=LINE_INDENT)
            sheet.cell(row=row, column=2, value=round(line.net_weight_kg, 3))
            sheet.cell(row=row, column=3, value=line.qty)
            sheet.cell(row=row, column=4, value=line.unit_price)
            sheet.cell(row=row, column=5, value=round(line.amount, 4))
            _number_formats(sheet, row)
            row += 1
    return row


def _number_formats(sheet: Worksheet, row: int) -> None:
    for column, number_format in ((2, WEIGHT_FORMAT), (3, QTY_FORMAT), (4, PRICE_FORMAT), (5, AMOUNT_FORMAT)):
        cell = sheet.cell(row=row, column=column)
        cell.number_format = number_format
        cell.alignment = Alignment(horizontal="center" if column in (2, 3, 4) else "right")


def _write_footer(sheet: Worksheet, doc: InvoiceDocument, symbol: str, trade_term: str, row: int) -> None:
    totals = compute_totals(doc)
    values = ("総合計 (Total)", totals.net_weight_kg, totals.qty, trade_term, totals.amount)
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=column, value=value)
        cell.font = BOLD
        cell.fill = TOTAL_FILL
        cell.border = Border(top=MEDIUM, bottom=MEDIUM, left=THIN, right=THIN)
        cell.alignment = Alignment(horizontal="center")
    sheet.cell(row=row, column=2).number_format = '0.0"kg"'
    sheet.cell(row=row, column=3).number_format = QTY_FORMAT
    sheet.cell(row=row, column=5).number_format = f'"{symbol}"#,##0.00'
    sheet.cell(row=row, column=5).alignment = Alignment(horizontal="right")

    summary = (
        ("個数（Number of pieces）", totals.pieces, "0"),
        ("総重量(Gross weight) Kg", totals.gross_weight_kg, '0.0"kg"'),
        ("原産国(Country of Origin)", "/".join(totals.origins), None),
    )
    for offset, (label, value, number_format) in enumerate(summary, start=2):
        sheet.cell(row=row + offset, column=1, value=label).font = BOLD
        cell = sheet.cell(row=row + offset, column=2, value=value)
        cell.font = BOLD
        cell.alignment = Alignment(horizontal="left")
        if number_format:
            cell.number_format = number_format
    sheet.cell(row=row + 2, column=4, value="署名(Signature)").font = BOLD
    for column in (4, 5):
        sheet.cell(row=row + 4, column=column).border = Border(bottom=THIN)


def full_invoice_to_xlsx(doc: InvoiceDocument, currency_label: str, currency_symbol: str, trade_term_label: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f"Invoice {doc.number}"[:31]
    table_start = _write_header(sheet, doc)
    footer_row = _write_table(sheet, doc, currency_label, table_start)
    _write_footer(sheet, doc, currency_symbol, trade_term_label, footer_row)
    for column, width in COLUMN_WIDTHS.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = sheet.cell(row=table_start + 1, column=1)
    sheet.print_title_rows = f"{table_start}:{table_start}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
