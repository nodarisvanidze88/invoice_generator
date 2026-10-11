"""Bulk product update from the Products Excel export: match by ID, then by JAN code; unmatched rows are created."""

from dataclasses import dataclass, field
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from pydantic import ValidationError
from sqlmodel import Session, select

from app.models import Product, ProductFields, utcnow
from app.services.excel_export import COLUMNS

NUMERIC_FIELDS = {"weight_g", "unit_price", "length_cm", "width_cm", "height_cm"}
BOOLEAN_FIELDS = {"keep_upright"}
TRUE_WORDS = {"yes", "y", "true", "1", "x", "✓"}
MAX_ERRORS = 20


class ProductImportError(ValueError):
    pass


@dataclass
class ProductImportResult:
    created: int = 0
    updated: int = 0
    skipped: int = 0
    errors: list[str] = field(default_factory=list)


def _clean(attr: str, value: Any) -> Any:
    if attr in BOOLEAN_FIELDS:
        return str(value).strip().lower() in TRUE_WORDS if value is not None else False
    if attr in NUMERIC_FIELDS:
        if value in (None, ""):
            return None
        try:
            return float(str(value).replace(",", "").strip())
        except ValueError as exc:
            raise ValueError(f"{attr} must be a number") from exc
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _read_rows(content: bytes) -> tuple[dict[int, str], list[tuple[int, tuple[Any, ...]]]]:
    try:
        sheet = load_workbook(BytesIO(content), read_only=True, data_only=True).active
    except Exception as exc:
        raise ProductImportError("Not a valid .xlsx file") from exc
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        raise ProductImportError("The sheet is empty")
    by_title = {title.lower(): attr for title, attr, _ in COLUMNS}
    columns = {i: by_title[str(h).strip().lower()] for i, h in enumerate(header) if h and str(h).strip().lower() in by_title}
    if "name" not in columns.values() and "id" not in columns.values():
        raise ProductImportError("Header row must contain 'ID' or 'Product Name' (use the Export Excel file as a template)")
    return columns, [(number, row) for number, row in enumerate(rows, start=2) if any(v not in (None, "") for v in row)]


def import_products(session: Session, content: bytes) -> ProductImportResult:
    columns, rows = _read_rows(content)
    products = list(session.exec(select(Product)))
    by_id = {p.id: p for p in products}
    by_code = {p.code.strip(): p for p in products if p.code.strip()}
    result = ProductImportResult()

    for number, row in rows:
        try:
            values = {attr: _clean(attr, row[i]) for i, attr in columns.items() if attr != "id" and i < len(row)}
            raw_id = row[next(i for i, a in columns.items() if a == "id")] if "id" in columns.values() else None
            product = by_id.get(int(raw_id)) if raw_id not in (None, "") else None
            product = product or by_code.get(str(values.get("code", "")).strip())
            if product is not None:
                merged = ProductFields.model_validate(product.model_dump(include=set(ProductFields.model_fields)) | values)
                product.sqlmodel_update(merged.model_dump())
                product.updated_at = utcnow()
                result.updated += 1
            else:
                product = Product.model_validate(ProductFields.model_validate(values).model_dump())
                result.created += 1
            session.add(product)
        except (ValueError, ValidationError) as exc:
            result.skipped += 1
            if len(result.errors) < MAX_ERRORS:
                message = exc.errors()[0]["msg"] if isinstance(exc, ValidationError) else str(exc)
                result.errors.append(f"Row {number}: {message}")
    session.commit()
    return result
