"""Detect where product data lives in an arbitrary workbook: which sheet, which header row, which column is which field."""

import re
from collections.abc import Iterator
from dataclasses import dataclass, field
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

HEADER_SCAN_ROWS = 15
SAMPLE_ROWS = 3
CODE_SAMPLE_ROWS = 20
CODE_COLUMN_MIN_SHARE = 0.6
MIN_PREFIX_ALIAS = 5
MIN_PREFIX_ALIAS_CJK = 2
NAME_BONUS = 5
FILL_SAMPLE_ROWS = 50
FILL_WEIGHT = 20
MAX_DATA_ROWS = 20_000
GRAMS_PER_KG = 1000

FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "id": ("id",),
    "code": ("jancode", "jan", "barcode", "ean", "eancode", "code", "sku", "itemcode", "productcode", "janコード", "商品コード", "バーコード"),
    "name": ("productname", "itemname", "name", "item", "description", "merchandise", "商品名", "品名", "商品"),
    "material": ("material", "materials", "素材", "材質"),
    "origin": ("origin", "countryoforigin", "madein", "coo", "原産国", "原産地"),
    "weight_g": ("weightg", "weight", "unitweight", "grams", "netweight", "weightkg", "重量", "重さ"),
    "unit_price": ("unitprice", "price", "cost", "unitcost", "単価", "価格"),
    "category": ("productcategory", "category", "カテゴリ", "カテゴリー", "分類"),
    "tariff_code": ("tariffcode", "hscode", "hs", "tariff", "関税コード"),
    "length_cm": ("lengthcm", "length", "長さ", "縦"),
    "width_cm": ("widthcm", "width", "幅", "横"),
    "height_cm": ("heightcm", "height", "高さ"),
    "keep_upright": ("keepupright", "upright", "thissideup", "天地無用"),
}
FIELD_ORDER = tuple(FIELD_ALIASES)
CODE_PATTERN = re.compile(r"^\d{8,14}(\s*/\s*\d{8,14})*$")
NORMALIZE_PATTERN = re.compile(r"[\s()（）\[\]/_\-.:#・]+")


class LayoutError(ValueError):
    pass


@dataclass(frozen=True)
class ColumnInfo:
    index: int
    letter: str
    header: str
    samples: list[str]


@dataclass
class SheetLayout:
    sheet: str
    header_row: int
    mapping: dict[str, int]
    weight_in_kg: bool = False


@dataclass
class SheetSummary:
    name: str
    score: int
    recognized: int
    header_row: int


@dataclass
class Inspection:
    sheets: list[SheetSummary]
    layout: SheetLayout
    columns: list[ColumnInfo]
    data_rows: int
    unmapped_headers: list[str] = field(default_factory=list)


def normalize(text: Any) -> str:
    return NORMALIZE_PATTERN.sub("", str(text or "").strip().lower())


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def open_workbook(content: bytes):
    try:
        return load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise LayoutError("Not a valid Excel file (.xlsx / .xlsm)") from exc


def _rows(sheet: ReadOnlyWorksheet, min_row: int = 1, max_row: int | None = None) -> Iterator[tuple[Any, ...]]:
    sheet.reset_dimensions()
    yield from sheet.iter_rows(min_row=min_row, max_row=max_row, values_only=True)


def _min_prefix(alias: str) -> int:
    return MIN_PREFIX_ALIAS if alias.isascii() else MIN_PREFIX_ALIAS_CJK


def match_headers(headers: list[str]) -> tuple[dict[str, int], bool]:
    """Assign each field the best column: exact alias beats prefix match; each column is used once."""
    candidates: list[tuple[int, int, str, int]] = []
    for index, header in enumerate(headers):
        normalized = normalize(header)
        if not normalized:
            continue
        for field_name, aliases in FIELD_ALIASES.items():
            if normalized in aliases:
                candidates.append((2, 0, field_name, index))
                continue
            prefix = max((len(a) for a in aliases if len(a) >= _min_prefix(a) and normalized.startswith(a)), default=0)
            if prefix:
                candidates.append((1, prefix, field_name, index))
    mapping: dict[str, int] = {}
    used: set[int] = set()
    for _quality, _length, field_name, index in sorted(candidates, key=lambda c: (-c[0], -c[1], c[3])):
        if field_name not in mapping and index not in used:
            mapping[field_name] = index
            used.add(index)
    weight_header = normalize(headers[mapping["weight_g"]]) if "weight_g" in mapping else ""
    return mapping, "kg" in weight_header


def _detect_code_column(sheet: ReadOnlyWorksheet, header_row: int, mapping: dict[str, int], width: int) -> None:
    if "code" in mapping:
        return
    samples = [row for row in _rows(sheet, header_row + 1, header_row + CODE_SAMPLE_ROWS) if any(v not in (None, "") for v in row)]
    used = set(mapping.values())
    for index in range(width):
        if index in used:
            continue
        values = [_cell_text(row[index]) for row in samples if index < len(row) and _cell_text(row[index])]
        if values and sum(bool(CODE_PATTERN.match(v)) for v in values) / len(values) >= CODE_COLUMN_MIN_SHARE:
            mapping["code"] = index
            return


def _best_header_row(sheet: ReadOnlyWorksheet) -> tuple[int, list[str], dict[str, int], bool]:
    best: tuple[int, int, list[str], dict[str, int], bool] = (-1, 1, [], {}, False)
    for number, row in enumerate(_rows(sheet, 1, HEADER_SCAN_ROWS), start=1):
        headers = [_cell_text(v) for v in row]
        mapping, in_kg = match_headers(headers)
        score = len(mapping) + (NAME_BONUS if "name" in mapping else 0)
        if score > best[0]:
            best = (score, number, headers, mapping, in_kg)
    return best[1], best[2], best[3], best[4]


def _fill_ratio(sheet: ReadOnlyWorksheet, header_row: int, mapping: dict[str, int]) -> float:
    """Share of rows below the header that really contain a product name (or code) — headers alone are not enough."""
    key = mapping.get("name", mapping.get("code"))
    if key is None:
        return 0.0
    rows = list(_rows(sheet, header_row + 1, header_row + FILL_SAMPLE_ROWS))
    if not rows:
        return 0.0
    return sum(1 for row in rows if key < len(row) and _cell_text(row[key])) / len(rows)


def _sheet_summary(sheet: ReadOnlyWorksheet) -> SheetSummary:
    header_row, _headers, mapping, _ = _best_header_row(sheet)
    fill = _fill_ratio(sheet, header_row, mapping)
    score = round(len(mapping) + (NAME_BONUS if "name" in mapping else 0) + FILL_WEIGHT * fill)
    return SheetSummary(name=sheet.title, score=score, recognized=len(mapping), header_row=header_row)


def describe_sheet(sheet: ReadOnlyWorksheet, header_row: int | None) -> tuple[SheetLayout, list[ColumnInfo], int]:
    if header_row is None:
        header_row, headers, mapping, in_kg = _best_header_row(sheet)
    else:
        header_values = next(_rows(sheet, header_row, header_row), ())
        headers = [_cell_text(v) for v in header_values]
        mapping, in_kg = match_headers(headers)

    data = [row for row in _rows(sheet, header_row + 1) if any(v not in (None, "") for v in row)]
    width = max([len(headers), *(len(row) for row in data[:SAMPLE_ROWS * 10])], default=0)
    _detect_code_column(sheet, header_row, mapping, width)
    columns = [
        ColumnInfo(
            index=index,
            letter=get_column_letter(index + 1),
            header=headers[index] if index < len(headers) else "",
            samples=[_cell_text(row[index])[:60] for row in data[:SAMPLE_ROWS] if index < len(row)],
        )
        for index in range(width)
    ]
    return SheetLayout(sheet=sheet.title, header_row=header_row, mapping=mapping, weight_in_kg=in_kg), columns, len(data)


def inspect_workbook(content: bytes, sheet_name: str | None = None, header_row: int | None = None) -> Inspection:
    workbook = open_workbook(content)
    worksheets = [ws for ws in workbook.worksheets if isinstance(ws, ReadOnlyWorksheet)]
    if not worksheets:
        raise LayoutError("The workbook has no sheets")
    summaries = [_sheet_summary(ws) for ws in worksheets]
    if sheet_name is None:
        chosen = worksheets[max(range(len(summaries)), key=lambda i: summaries[i].score)]
    else:
        chosen = next((ws for ws in worksheets if ws.title == sheet_name), None)
        if chosen is None:
            raise LayoutError(f"Sheet '{sheet_name}' not found")
    layout, columns, data_rows = describe_sheet(chosen, header_row)
    mapped = set(layout.mapping.values())
    unmapped = [f"{c.letter} · {c.header}" for c in columns if c.index not in mapped and c.header]
    return Inspection(sheets=summaries, layout=layout, columns=columns, data_rows=data_rows, unmapped_headers=unmapped)


def read_table(content: bytes, layout: SheetLayout) -> list[tuple[int, dict[str, Any]]]:
    """Rows below the header as {field: raw value}; weight converted to grams when the column is in kg."""
    if "name" not in layout.mapping and "id" not in layout.mapping:
        raise LayoutError("Map at least the Product Name column (or ID) before importing")
    workbook = open_workbook(content)
    if layout.sheet not in workbook.sheetnames:
        raise LayoutError(f"Sheet '{layout.sheet}' not found")
    sheet = workbook[layout.sheet]
    result: list[tuple[int, dict[str, Any]]] = []
    for number, row in enumerate(_rows(sheet, layout.header_row + 1), start=layout.header_row + 1):
        values = {name: row[index] if index < len(row) else None for name, index in layout.mapping.items()}
        if all(v in (None, "") for v in values.values()):
            continue
        if layout.weight_in_kg and isinstance(values.get("weight_g"), (int, float)):
            values["weight_g"] = values["weight_g"] * GRAMS_PER_KG
        result.append((number, values))
        if len(result) > MAX_DATA_ROWS:
            raise LayoutError(f"Too many rows — the limit is {MAX_DATA_ROWS:,} per import")
    return result
